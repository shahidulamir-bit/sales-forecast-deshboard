from pathlib import Path
from openpyxl import load_workbook

p = Path(r'd:\streamlit webapp\active forcast.xlsx')
print('exists=', p.exists())
print('size=', p.stat().st_size if p.exists() else 'missing')
if p.exists():
    wb = load_workbook(p, read_only=True)
    print('sheetnames=', wb.sheetnames)
    for name in wb.sheetnames:
        ws = wb[name]
        print('sheet=', name, 'rows=', ws.max_row, 'cols=', ws.max_column)
        for row in ws.iter_rows(min_row=1, max_row=min(3, ws.max_row), values_only=True):
            print(row)
